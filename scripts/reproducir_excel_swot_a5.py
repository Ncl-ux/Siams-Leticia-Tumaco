"""SIAMS A5: reconstruye UN Excel SWOT desde exportaciones HydroWeb originales.

Entradas: data/raw/swot/<sede>/plot_*.csv, sin lectura de data/processed/.
Salida: 1 XLSX en data/processed/swot/ con datos, calidad exploratoria y procedencia.

No certifica banderas oficiales RiverSP ni sustituye los productos ya publicados.
Requiere pandas, numpy y openpyxl (incluidos en requirements.txt de SIAMS).

Desde la raíz del proyecto:
    python scripts/reproducir_excel_swot_a5.py
    python scripts/reproducir_excel_swot_a5.py --salida data/processed/swot/mi_version.xlsx
"""
from __future__ import annotations

import argparse
from hashlib import sha256
from pathlib import Path
from tempfile import TemporaryDirectory
import sys

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

from procesar_swot import regenerar, RAW as RAW_PREDETERMINADO  # noqa: E402

SALIDA_PREDETERMINADA = ROOT / 'data' / 'processed' / 'swot' / 'SWOT_A5_reproducido_desde_plot.xlsx'

DICCIONARIO = [
    ('fecha_utc', 'Fecha de adquisición, UTC', 'ISO 8601', 'Fecha declarada por HydroWeb'),
    ('sede', 'Sede de agrupación de exportaciones', 'texto', 'Carpeta de la entrada original'),
    ('reach_id', 'Identificador del tramo fluvial', 'identificador', 'Identificador de HydroWeb'),
    ('wse_m', 'Elevación de superficie del agua', 'm', 'HydroWeb SWOT, no validada con banderas RiverSP'),
    ('wse_incertidumbre_m', 'Incertidumbre reportada de WSE', 'm', 'Exportación plot original'),
    ('ancho_m', 'Ancho estimado del río', 'm', 'Exportación plot original'),
    ('ancho_incertidumbre_m', 'Incertidumbre de ancho', 'm', 'Exportación plot original'),
    ('area_m2', 'Área superficial estimada del agua', 'm²', 'Exportación plot original'),
    ('area_incertidumbre_m2', 'Incertidumbre del área', 'm²', 'Exportación plot original'),
    ('pendiente_m_m', 'Pendiente superficial estimada', 'm/m', 'Exportación plot original'),
    ('pendiente_incertidumbre_m_m', 'Incertidumbre de pendiente', 'm/m', 'Exportación plot original'),
    ('salto_absoluto_wse_m', 'Diferencia absoluta de WSE entre observaciones sucesivas', 'm', 'Calculada; no implica error'),
    ('umbral_robusto_salto_wse_m', 'Umbral mediana + 6×1,4826×MAD', 'm', 'Control exploratorio, no oficial'),
    ('revisar_salto_wse', 'Alerta estadística sobre un salto de WSE', 'SI/NO', 'No sustituye control oficial RiverSP'),
    ('bandera_calidad_oficial', 'Disponibilidad de bandera de calidad oficial', 'texto', 'NO_PRESENTE_EN_CSV'),
]


def _sha256(ruta: Path) -> str:
    h = sha256()
    with ruta.open('rb') as f:
        for bloque in iter(lambda: f.read(1024 * 1024), b''):
            h.update(bloque)
    return h.hexdigest()


def _ajustar_hoja(hoja, *, tabla_ancha=False):
    hoja.freeze_panes = 'D2' if tabla_ancha else 'A2'
    hoja.auto_filter.ref = hoja.dimensions
    hoja.sheet_view.showGridLines = False
    hoja.row_dimensions[1].height = 27
    fin = hoja.max_column
    for col in hoja[1]:
        col.fill = PatternFill('solid', fgColor='16435B')
        col.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        col.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for c in range(1, fin + 1):
        cabecera = str(hoja.cell(1, c).value or '')
        largo = min(40, max(14, len(cabecera) + 4))
        if 'archivo' in cabecera or 'origen' in cabecera or 'sha256' in cabecera:
            largo = 48
        elif cabecera in ('fecha_utc', 'variable', 'descripcion', 'nota', 'limitacion'):
            largo = 34
        hoja.column_dimensions[get_column_letter(c)].width = largo
    for i in range(2, hoja.max_row + 1):
        for j in range(1, fin + 1):
            celda = hoja.cell(i, j)
            if i % 2 == 0:
                celda.fill = PatternFill('solid', fgColor='EEF5F8')
            if isinstance(celda.value, float):
                celda.number_format = '0.000000'
            celda.alignment = Alignment(vertical='center')


def construir_excel(*, raw: Path = RAW_PREDETERMINADO, salida: Path = SALIDA_PREDETERMINADA) -> dict:
    """Regenera desde raw, no desde los CSV ya publicados; aborta sin sobrescribir."""
    raw = Path(raw).resolve()
    salida = Path(salida).resolve()
    if salida.exists():
        raise FileExistsError(f'El Excel ya existe y no se sobrescribirá: {salida}')
    if not raw.is_dir():
        raise FileNotFoundError(f'Faltan las exportaciones originales de SWOT: {raw}')
    fuentes = sorted(raw.rglob('plot_*.csv'))
    if not fuentes:
        raise FileNotFoundError(f'No hay ningún plot_*.csv dentro de {raw}')

    with TemporaryDirectory(prefix='siams_a5_swot_') as tmp:
        carpeta_tmp = Path(tmp)
        resumen = regenerar(raw, carpeta_tmp)
        datasets = []
        for ruta in sorted(carpeta_tmp.glob('SWOT_*_reach_*.csv')):
            tabla = pd.read_csv(ruta, dtype={'reach_id': str}, encoding='utf-8-sig')
            # La sede deriva del nombre de archivo generado por el procesador.
            sede = ruta.name.removeprefix('SWOT_').split('_reach_')[0]
            tabla.insert(0, 'sede', sede)
            datasets.append(tabla)
        if len(datasets) != len(resumen):
            raise RuntimeError('La cantidad de tramos regenerados y resumidos no coincide')
        registros = pd.concat(datasets, ignore_index=True)
        columnas = ['sede', 'reach_id', 'fecha_utc']
        registros = registros.sort_values(columnas).reset_index(drop=True)
        if len(registros) != int(resumen['fechas_unicas'].sum()):
            raise RuntimeError('La cantidad de filas no coincide con el resumen')

        manifiesto = pd.DataFrame([
            {'ruta_en_raw': p.relative_to(raw).as_posix(),
             'bytes': p.stat().st_size,
             'sha256': _sha256(p),
             'fuente_declarada': 'HydroWeb / SWOT; exportación plot_ID recibida'}
            for p in fuentes
        ])
        diccionario = pd.DataFrame(DICCIONARIO, columns=['variable', 'descripcion', 'unidad', 'metodologia_y_limitacion'])
        metodologia = pd.DataFrame([
            ('Propósito', 'Reproducibilidad A5, generación de un Excel desde datos originales'),
            ('Entrada', 'data/raw/swot/*/plot_*.csv sin modificaciones'),
            ('Proceso', 'scripts/procesar_swot.py + scripts/reproducir_excel_swot_a5.py'),
            ('Reaches', str(len(resumen))),
            ('Registros', str(len(registros))),
            ('Archivos originales', str(len(fuentes))),
            ('Alertas exploratorias', str(int(resumen.saltos_para_revisar.sum()))),
            ('Control de calidad', 'Umbral estadístico; NO equivale a banderas oficiales RiverSP'),
            ('Asignación Mira/Rosario', 'Pendiente de verificar contra los metadatos originales; no inferir por prefijo'),
            ('Unión temporal', 'Por fecha UTC con tolerancia de 3 segundos entre variables'),
            ('Trazabilidad', 'Cada archivo de entrada tiene ruta, tamaño en bytes y SHA-256 en Fuentes_raw'),
            ('Archivo histórico', 'NO sobrescribe ni certifica equivalencia con los Excel de otras fuentes'),
            ('Fecha de emisión', 'No se incluye para evitar cambiar el contenido entre ejecuciones'),
        ], columns=['campo', 'valor'])
        # No escribir a un archivo final parcial ante un fallo de ExcelWriter.
        salida.parent.mkdir(parents=True, exist_ok=True)
        temporal = salida.with_name(salida.stem + '.escribiendo.xlsx')
        try:
            if temporal.exists():
                raise FileExistsError(f'Hay un archivo temporal previo: {temporal}')
            with pd.ExcelWriter(temporal, engine='openpyxl') as writer:
                hojas = {
                    'Resumen': resumen,
                    'Observaciones': registros,
                    'Fuentes_raw': manifiesto,
                    'Diccionario': diccionario,
                    'Metodologia': metodologia,
                }
                for nombre, df in hojas.items():
                    df.to_excel(writer, sheet_name=nombre, index=False)
                    _ajustar_hoja(writer.sheets[nombre], tabla_ancha=(nombre == 'Observaciones'))
                ws = writer.sheets['Resumen']
                # Color de advertencia en las alertas estadísticas (sin borrar datos).
                col_alertas = list(resumen.columns).index('saltos_para_revisar') + 1
                letra = get_column_letter(col_alertas)
                ws.conditional_formatting.add(f'{letra}2:{letra}{ws.max_row}',
                    CellIsRule(operator='greaterThan', formula=['0'],
                        fill=PatternFill('solid', fgColor='FCE8D5')))
            temporal.replace(salida)
        finally:
            if temporal.exists():
                temporal.unlink()
        return {'ruta': str(salida), 'plot_originales': len(fuentes),
                'tramos': len(resumen), 'observaciones': len(registros),
                'alertas_exploratorias': int(resumen.saltos_para_revisar.sum())}


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--raw', type=Path, default=RAW_PREDETERMINADO)
    p.add_argument('--salida', type=Path, default=SALIDA_PREDETERMINADA)
    args = p.parse_args()
    resultado = construir_excel(raw=args.raw, salida=args.salida)
    for clave, valor in resultado.items():
        print(f'{clave}: {valor}')


if __name__ == '__main__':
    main()
